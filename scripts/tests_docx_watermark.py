"""
Tests for KadhaScript — Screenplay DOCX Export Custom Watermark (Phase 4).

Verifies:
1. Empty or whitespace-only watermark produces no watermark markup or header artifact.
2. English watermark (e.g. CONFIDENTIAL) is properly added to header VML shape.
3. Malayalam watermark (e.g. കരട് പ്രതി) preserves Unicode characters and uses Noto Sans Malayalam font.
4. Special characters (e.g. Draft & Review <2026> - "Confidential") are properly XML-escaped and valid.
5. Multi-page screenplay documents configure headers that repeat across pages.
6. Screenplay content (scene headings, action, character, dialogue) remains completely intact.
7. Existing margins, breaks, and header/footer behavior are preserved.
8. No body layout pollution (watermark does not add ordinary body paragraphs or alter document flow).
"""

import io
from docx import Document
from django.test import TestCase
from django.contrib.auth import get_user_model

from scripts.models import Script, Scene, ScriptElement
from scripts.services.docx_export import generate_screenplay_docx, _apply_docx_watermark

User = get_user_model()


class DOCXExportWatermarkTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            username='docx_wm_user',
            email='docx_wm@example.com',
            password='Password123!'
        )

    def test_empty_or_whitespace_watermark_adds_no_watermark(self):
        """When watermark is empty or whitespace-only, no VML watermark shape is added."""
        for wm_val in ['', '   ', None]:
            script = Script.objects.create(
                user=self.user,
                title='Clean Unwatermarked DOCX',
                language='Malayalam',
                watermark=wm_val or ''
            )
            sc = Scene.objects.create(
                script=script,
                scene_number=1,
                heading='INT. OFFICE - DAY',
                order=0
            )
            ScriptElement.objects.create(
                scene=sc,
                element_type='action',
                content='Normal screenplay action line.',
                order=0
            )

            docx_bytes = generate_screenplay_docx(script)
            doc = Document(io.BytesIO(docx_bytes))

            # Header must not contain PowerPlusWaterMarkObject
            header_xml = ""
            for s in doc.sections:
                for p in s.header.paragraphs:
                    header_xml += p._p.xml

            self.assertNotIn('PowerPlusWaterMarkObject', header_xml)
            self.assertNotIn('v:shape', header_xml)
            self.assertNotIn('v:textpath', header_xml)

            # Body paragraphs must contain only the screenplay content
            body_text = " ".join(p.text for p in doc.paragraphs)
            self.assertIn('INT. OFFICE - DAY', body_text)
            self.assertIn('Normal screenplay action line.', body_text)
            script.delete()

    def test_english_watermark_included_in_docx(self):
        """English watermark is included in the header VML shape without polluting body flow."""
        script = Script.objects.create(
            user=self.user,
            title='Confidential DOCX',
            language='Malayalam',
            watermark='CONFIDENTIAL'
        )
        sc = Scene.objects.create(
            script=script,
            scene_number=1,
            heading='INT. STUDY - DAY',
            order=0
        )
        ScriptElement.objects.create(
            scene=sc,
            element_type='action',
            content='Secret blueprints rest on the mahogany desk.',
            order=0
        )

        docx_bytes = generate_screenplay_docx(script)
        doc = Document(io.BytesIO(docx_bytes))

        # 1. Header has VML watermark object
        header_xml = "".join(p._p.xml for s in doc.sections for p in s.header.paragraphs)
        self.assertIn('PowerPlusWaterMarkObject', header_xml)
        self.assertIn('string="CONFIDENTIAL"', header_xml)
        self.assertIn('rotation:-35', header_xml)

        # 2. Body paragraphs do not contain watermark text as normal paragraphs
        body_texts = [p.text for p in doc.paragraphs]
        self.assertNotIn('CONFIDENTIAL', body_texts)
        full_body = " ".join(body_texts)
        self.assertIn('Secret blueprints rest on the mahogany desk.', full_body)
        script.delete()

    def test_malayalam_watermark_preserves_unicode_and_font(self):
        """Malayalam watermark preserves Unicode characters and specifies Noto Sans Malayalam."""
        script = Script.objects.create(
            user=self.user,
            title='മലയാളം തിരക്കഥ DOCX',
            language='Malayalam',
            watermark='കരട് പ്രതി'
        )
        sc = Scene.objects.create(
            script=script,
            scene_number=1,
            heading='INT. വീട് - DAY',
            order=0
        )
        ScriptElement.objects.create(
            scene=sc,
            element_type='action',
            content='നായകൻ അകത്തേക്ക് വരുന്നു.',
            order=0
        )

        docx_bytes = generate_screenplay_docx(script)
        doc = Document(io.BytesIO(docx_bytes))

        header_xml = "".join(p._p.xml for s in doc.sections for p in s.header.paragraphs)
        self.assertIn('PowerPlusWaterMarkObject', header_xml)
        self.assertIn('string="കരട് പ്രതി"', header_xml)
        self.assertIn('Noto Sans Malayalam', header_xml)

        # Body remains intact
        full_body = " ".join(p.text for p in doc.paragraphs)
        self.assertIn('INT. വീട് - DAY', full_body)
        self.assertIn('നായകൻ അകത്തേക്ക് വരുന്നു.', full_body)
        script.delete()

    def test_special_characters_escaped_safely(self):
        """Special XML characters in watermark are safely escaped without invalidating DOCX."""
        script = Script.objects.create(
            user=self.user,
            title='Special Characters DOCX',
            language='Malayalam',
            watermark='Draft & Review <2026> - "Confidential"'
        )
        sc = Scene.objects.create(
            script=script,
            scene_number=1,
            heading='EXT. PARK - DAY',
            order=0
        )
        ScriptElement.objects.create(
            scene=sc,
            element_type='action',
            content='Children play in the park.',
            order=0
        )

        docx_bytes = generate_screenplay_docx(script)
        # Must be valid DOCX parseable by python-docx
        doc = Document(io.BytesIO(docx_bytes))

        header_xml = "".join(p._p.xml for s in doc.sections for p in s.header.paragraphs)
        self.assertIn('PowerPlusWaterMarkObject', header_xml)
        self.assertIn('Draft &amp; Review &lt;2026&gt; - &quot;Confidential&quot;', header_xml)
        script.delete()

    def test_multi_page_document_repeats_watermark(self):
        """Multi-page screenplay has watermark in section header that repeats across pages."""
        script = Script.objects.create(
            user=self.user,
            title='Multipage DOCX',
            language='Malayalam',
            watermark='FOR INTERNAL REVIEW'
        )
        for i in range(1, 4):
            sc = Scene.objects.create(
                script=script,
                scene_number=i,
                heading=f'INT. ROOM {i} - DAY',
                order=i - 1
            )
            ScriptElement.objects.create(
                scene=sc,
                element_type='action',
                content=f'Action line for scene {i}. ' * 8,
                order=0
            )

        docx_bytes = generate_screenplay_docx(script)
        doc = Document(io.BytesIO(docx_bytes))

        # Verify page breaks exist
        page_break_count = 0
        for p in doc.paragraphs:
            for r in p.runs:
                if 'w:br' in r._r.xml and 'type="page"' in r._r.xml:
                    page_break_count += 1
        self.assertGreaterEqual(page_break_count, 3)

        # Verify every section header contains the watermark
        for section in doc.sections:
            h_xml = "".join(p._p.xml for p in section.header.paragraphs)
            self.assertIn('PowerPlusWaterMarkObject', h_xml)
            self.assertIn('string="FOR INTERNAL REVIEW"', h_xml)
        script.delete()

    def test_existing_screenplay_elements_intact(self):
        """All screenplay element types (scene_heading, action, character, dialogue, parenthetical, transition, shot) remain intact."""
        script = Script.objects.create(
            user=self.user,
            title='Full Element Script',
            language='Malayalam',
            watermark='DRAFT 2'
        )
        sc = Scene.objects.create(script=script, scene_number=1, heading='INT. COURT - DAY', order=0)
        ScriptElement.objects.create(scene=sc, element_type='action', content='The judge enters.', order=0)
        ScriptElement.objects.create(scene=sc, element_type='character', content='JUDGE', order=1)
        ScriptElement.objects.create(scene=sc, element_type='parenthetical', content='calmly', order=2)
        ScriptElement.objects.create(scene=sc, element_type='dialogue', content='Court is now in session.', order=3)
        ScriptElement.objects.create(scene=sc, element_type='shot', content='WIDE ON JURY', order=4)
        ScriptElement.objects.create(scene=sc, element_type='transition', content='FADE OUT', order=5)

        docx_bytes = generate_screenplay_docx(script)
        doc = Document(io.BytesIO(docx_bytes))
        full_text = " ".join(p.text for p in doc.paragraphs)

        self.assertIn('INT. COURT - DAY', full_text)
        self.assertIn('The judge enters.', full_text)
        self.assertIn('JUDGE', full_text)
        self.assertIn('(calmly)', full_text)
        self.assertIn('Court is now in session.', full_text)
        self.assertIn('WIDE ON JURY', full_text)
        self.assertIn('FADE OUT', full_text)
        script.delete()

    def test_no_body_layout_pollution(self):
        """Adding a watermark does not add any extra body paragraphs or change body paragraph count."""
        script_without = Script.objects.create(
            user=self.user,
            title='Baseline Script',
            language='Malayalam',
            watermark=''
        )
        sc1 = Scene.objects.create(script=script_without, scene_number=1, heading='INT. CAFE - DAY', order=0)
        ScriptElement.objects.create(scene=sc1, element_type='action', content='Coffee aroma fills the air.', order=0)

        docx_bytes_without = generate_screenplay_docx(script_without)
        doc_without = Document(io.BytesIO(docx_bytes_without))
        para_count_without = len(doc_without.paragraphs)

        # Now test with watermark
        script_without.watermark = 'CONFIDENTIAL WATERMARK'
        script_without.save()

        docx_bytes_with = generate_screenplay_docx(script_without)
        doc_with = Document(io.BytesIO(docx_bytes_with))
        para_count_with = len(doc_with.paragraphs)

        # Body paragraph count must be identical
        self.assertEqual(para_count_with, para_count_without)

        # Body paragraph texts must be identical
        self.assertEqual(
            [p.text for p in doc_with.paragraphs],
            [p.text for p in doc_without.paragraphs]
        )
        script_without.delete()

    def test_invalid_xml_control_characters_stripped_safely(self):
        """XML 1.0-invalid control characters are stripped while preserving surrounding text and escaping."""
        script = Script.objects.create(
            user=self.user,
            title='Control Character Hardening DOCX',
            language='Malayalam',
            watermark=''
        )
        # Assign programmatic value containing NUL (\x00) and BEL (\x07) control characters
        script.watermark = 'Draft\x00Review\x07<2026>'
        sc = Scene.objects.create(
            script=script,
            scene_number=1,
            heading='INT. LAB - NIGHT',
            order=0
        )
        ScriptElement.objects.create(
            scene=sc,
            element_type='action',
            content='Testing XML control character sanitization.',
            order=0
        )

        # 1 & 2. DOCX generation does not crash and resulting DOCX is valid
        docx_bytes = generate_screenplay_docx(script)
        doc = Document(io.BytesIO(docx_bytes))

        header_xml = "".join(p._p.xml for s in doc.sections for p in s.header.paragraphs)
        self.assertIn('PowerPlusWaterMarkObject', header_xml)

        # 3, 4 & 5. Valid surrounding text remains, invalid control chars are absent, and special-char escaping works
        self.assertIn('string="DraftReview&lt;2026&gt;"', header_xml)
        self.assertNotIn('\x00', header_xml)
        self.assertNotIn('\x07', header_xml)
        script.delete()
