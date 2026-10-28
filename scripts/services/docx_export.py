import io
import html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import parse_xml

DEFAULT_MALAYALAM_FONT = "Noto Sans Malayalam"

def _set_run_font(run, font_name=DEFAULT_MALAYALAM_FONT, size=None, bold=None, italic=None, color=None):
    """
    Sets font properties and explicit OpenXML font tags (w:ascii, w:hAnsi, w:cs).
    Setting w:cs ensures Microsoft Word and compatible viewers render Malayalam
    complex scripts (conjuncts, vowel signs, chillus) with the designated font.
    """
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _apply_docx_watermark(doc, watermark_text: str):
    """
    Applies a standard Word VML diagonal watermark (PowerPlusWaterMarkObject)
    to all sections in the document via the header.
    Runs behind document text without affecting body flow, margins, or pagination.
    """
    if not watermark_text:
        return

    # Remove XML 1.0-invalid control characters before XML construction
    sanitized_watermark = ''.join(
        ch for ch in watermark_text
        if (
            (cp := ord(ch)) in (0x09, 0x0A, 0x0D)
            or (0x20 <= cp <= 0xD7FF)
            or (0xE000 <= cp <= 0xFFFD)
            or (0x10000 <= cp <= 0x10FFFF)
        )
    )
    clean_watermark = ' '.join(sanitized_watermark.split())
    if not clean_watermark:
        return

    # Safely escape for XML attributes, including quotes and special characters
    escaped_text = html.escape(clean_watermark, quote=True)

    # Check if watermark contains Malayalam characters
    has_malayalam = any(0x0D00 <= ord(ch) <= 0x0D7F for ch in clean_watermark)
    font_family = f"&quot;{DEFAULT_MALAYALAM_FONT}&quot;,Arial,sans-serif" if has_malayalam else "Arial,sans-serif"

    # Standard Word VML diagonal watermark shape definition
    vml_xml = (
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:o="urn:schemas-microsoft-com:office:office">'
        '  <w:pPr>'
        '    <w:pStyle w:val="Header"/>'
        '  </w:pPr>'
        '  <w:r>'
        '    <w:rPr>'
        '      <w:noProof/>'
        '    </w:rPr>'
        '    <w:pict>'
        '      <v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">'
        '        <v:path textpathok="t" o:connecttype="rect"/>'
        '        <v:textpath on="t" fitshape="t"/>'
        '        <o:lock v:ext="edit" text="t" shapetype="t"/>'
        '      </v:shapetype>'
        '      <v:shape id="PowerPlusWaterMarkObject" o:spid="_x0000_s1025" type="#_x0000_t136" '
        '        style="position:absolute;left:0;text-align:left;margin-left:0;margin-top:0;width:520pt;height:120pt;rotation:-35;z-index:-251658752;mso-position-horizontal:center;mso-position-horizontal-relative:page;mso-position-vertical:center;mso-position-vertical-relative:page" '
        '        fillcolor="#d0d0d0" stroked="f">'
        '        <v:fill opacity="0.15"/>'
        f'        <v:textpath style="font-family:{font_family};font-size:1pt" string="{escaped_text}"/>'
        '      </v:shape>'
        '    </w:pict>'
        '  </w:r>'
        '</w:p>'
    )

    for section in doc.sections:
        headers = [section.header]
        if getattr(section, 'different_first_page_header_footer', False):
            headers.append(section.first_page_header)

        for hdr in headers:
            new_p = parse_xml(vml_xml)
            if len(hdr.paragraphs) == 1 and not hdr.paragraphs[0].text.strip():
                hdr.paragraphs[0]._p.getparent().replace(hdr.paragraphs[0]._p, new_p)
            else:
                hdr._element.append(new_p)


def generate_screenplay_docx(script, include_notes=False) -> bytes:
    """Generates a screenplay formatted Microsoft Word document (DOCX) with full Malayalam Unicode support."""
    doc = Document()

    # Configure default Normal style with Malayalam font and complex script properties
    style = doc.styles['Normal']
    style.font.name = DEFAULT_MALAYALAM_FONT
    style.font.size = Pt(11)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), DEFAULT_MALAYALAM_FONT)
    rFonts.set(qn('w:hAnsi'), DEFAULT_MALAYALAM_FONT)
    rFonts.set(qn('w:cs'), DEFAULT_MALAYALAM_FONT)

    # Standard Screenplay Margins (1.5" Left, 1.0" Right, 1.0" Top, 1.0" Bottom)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

    # Document Title Block (Cover Page)
    title_page = getattr(script, 'title_page', None)
    effective_title = title_page.get_effective_title() if title_page else script.title
    effective_author = title_page.get_effective_author() if title_page else (
        script.author_name or (script.user.profile.pen_name if hasattr(script.user, 'profile') else (script.user.get_full_name() or script.user.username))
    )

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(72)
    title_p.paragraph_format.space_after = Pt(12)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(effective_title.upper())
    _set_run_font(title_run, size=Pt(22), bold=True)

    if title_page and title_page.subtitle:
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_before = Pt(0)
        sub_p.paragraph_format.space_after = Pt(12)
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(title_page.subtitle.strip())
        _set_run_font(sub_run, size=Pt(12), italic=True, color=RGBColor(80, 80, 80))

    by_p = doc.add_paragraph()
    by_p.paragraph_format.space_before = Pt(12)
    by_p.paragraph_format.space_after = Pt(6)
    by_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    by_run = by_p.add_run("എഴുതിയത് / Written by\n")
    _set_run_font(by_run, size=Pt(11))
    author_run = by_p.add_run(effective_author)
    _set_run_font(author_run, size=Pt(14), bold=True)

    if title_page and title_page.adaptation_credits:
        cred_p = doc.add_paragraph()
        cred_p.paragraph_format.space_before = Pt(8)
        cred_p.paragraph_format.space_after = Pt(18)
        cred_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cred_run = cred_p.add_run(title_page.adaptation_credits.strip())
        _set_run_font(cred_run, size=Pt(10), color=RGBColor(70, 70, 70))

    has_bottom_metadata = title_page and (
        title_page.draft_revision or title_page.draft_date or
        title_page.copyright_registration or
        title_page.contact_name or title_page.contact_email or title_page.contact_phone
    )

    if has_bottom_metadata:
        meta_p = doc.add_paragraph()
        meta_p.paragraph_format.space_before = Pt(72)
        meta_p.paragraph_format.space_after = Pt(12)
        meta_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        bottom_parts = []
        if title_page.draft_revision:
            bottom_parts.append(title_page.draft_revision.strip())
        if title_page.draft_date:
            bottom_parts.append(title_page.draft_date.strip())
        if title_page.copyright_registration:
            bottom_parts.append(title_page.copyright_registration.strip())
        
        contact_parts = []
        if title_page.contact_name:
            contact_parts.append(f"Contact: {title_page.contact_name.strip()}")
        if title_page.contact_email:
            contact_parts.append(title_page.contact_email.strip())
        if title_page.contact_phone:
            contact_parts.append(title_page.contact_phone.strip())

        full_footer_text = "\n".join(bottom_parts)
        if contact_parts:
            full_footer_text += ("\n\n" if full_footer_text else "") + "\n".join(contact_parts)

        meta_run = meta_p.add_run(full_footer_text)
        _set_run_font(meta_run, size=Pt(9.5), color=RGBColor(90, 90, 90))
    else:
        meta_p = doc.add_paragraph()
        meta_p.paragraph_format.space_before = Pt(36)
        meta_p.paragraph_format.space_after = Pt(12)
        meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta_p.add_run(
            f"Format: {script.script_type}  |  Genre: {script.genre}  |  Language: {script.language}\n"
            f"KadhaScript Screenplay Management Platform"
        )
        _set_run_font(meta_run, size=Pt(9.5), color=RGBColor(120, 120, 120))

    doc.add_page_break()

    # Scenes and Elements
    scenes = script.get_ordered_scenes()

    for idx, scene in enumerate(scenes):
        if idx > 0:
            doc.add_page_break()
        # Scene Heading from Scene Model
        scene_p = doc.add_paragraph()
        scene_p.paragraph_format.space_before = Pt(16)
        scene_p.paragraph_format.space_after = Pt(6)
        scene_p.paragraph_format.keep_with_next = True
        if scene.is_sequence:
            heading_text = scene.full_display_heading.upper()
        else:
            heading_text = f"{scene.scene_identifier} : {scene.clean_heading.upper()}"
        scene_run = scene_p.add_run(heading_text)
        _set_run_font(scene_run, size=Pt(11), bold=True)

        elements = scene.elements.all().order_by('order', 'id')
        for elem in elements:
            text = (elem.content or '').strip()
            if not text:
                continue

            if elem.element_type == 'scene_heading':
                clean_heading_upper = scene.clean_heading.strip().upper()
                raw_heading_upper = (scene.heading or '').strip().upper()
                elem_text_upper = text.strip().upper()
                if elem.order == 0 and (
                    elem_text_upper == clean_heading_upper
                    or elem_text_upper == raw_heading_upper
                ):
                    continue
                elem_p = doc.add_paragraph()
                elem_p.paragraph_format.space_before = Pt(14)
                elem_p.paragraph_format.space_after = Pt(6)
                elem_p.paragraph_format.keep_with_next = True
                run = elem_p.add_run(text.upper())
                _set_run_font(run, size=Pt(11), bold=True)

            elif elem.element_type == 'action':
                # Handle paragraphs separated by blank lines
                raw_paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
                if not raw_paragraphs:
                    raw_paragraphs = [text]
                for p_idx, para in enumerate(raw_paragraphs):
                    elem_p = doc.add_paragraph()
                    elem_p.paragraph_format.space_before = Pt(4) if p_idx == 0 else Pt(2)
                    elem_p.paragraph_format.space_after = Pt(6)
                    elem_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    run = elem_p.add_run(para)
                    _set_run_font(run, size=Pt(11))

            elif elem.element_type == 'character':
                elem_p = doc.add_paragraph()
                elem_p.paragraph_format.left_indent = Inches(2.2)
                elem_p.paragraph_format.space_before = Pt(12)
                elem_p.paragraph_format.space_after = Pt(2)
                elem_p.paragraph_format.keep_with_next = True
                run = elem_p.add_run(text.upper())
                _set_run_font(run, size=Pt(11), bold=True)

            elif elem.element_type == 'dialogue':
                elem_p = doc.add_paragraph()
                elem_p.paragraph_format.left_indent = Inches(1.2)
                elem_p.paragraph_format.right_indent = Inches(1.2)
                elem_p.paragraph_format.space_before = Pt(0)
                elem_p.paragraph_format.space_after = Pt(6)
                run = elem_p.add_run(text)
                _set_run_font(run, size=Pt(11))

            elif elem.element_type == 'parenthetical':
                elem_p = doc.add_paragraph()
                elem_p.paragraph_format.left_indent = Inches(1.7)
                elem_p.paragraph_format.right_indent = Inches(1.7)
                elem_p.paragraph_format.space_before = Pt(0)
                elem_p.paragraph_format.space_after = Pt(2)
                elem_p.paragraph_format.keep_with_next = True
                clean_text = text if (text.startswith('(') and text.endswith(')')) else f"({text})"
                run = elem_p.add_run(clean_text)
                _set_run_font(run, size=Pt(10))

            elif elem.element_type == 'transition':
                elem_p = doc.add_paragraph()
                elem_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elem_p.paragraph_format.space_before = Pt(12)
                elem_p.paragraph_format.space_after = Pt(12)
                run = elem_p.add_run(text.upper())
                _set_run_font(run, size=Pt(11), bold=True)

            elif elem.element_type == 'shot':
                if elem.order == 0 and scene.is_sequence and text.upper() == heading_text.upper():
                    continue
                elem_p = doc.add_paragraph()
                elem_p.paragraph_format.space_before = Pt(8)
                elem_p.paragraph_format.space_after = Pt(4)
                run = elem_p.add_run(text.upper())
                _set_run_font(run, size=Pt(11), bold=True)

            elif elem.element_type == 'note' and include_notes:
                raw_paragraphs = [p.strip() for p in text.split('\n\n') if p.strip()]
                if not raw_paragraphs:
                    raw_paragraphs = [text]
                for p_idx, para in enumerate(raw_paragraphs):
                    elem_p = doc.add_paragraph()
                    elem_p.paragraph_format.left_indent = Inches(0.5)
                    elem_p.paragraph_format.space_before = Pt(4) if p_idx == 0 else Pt(2)
                    elem_p.paragraph_format.space_after = Pt(6)
                    elem_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    note_text = para if (para.startswith('[[') and para.endswith(']]')) else f"[[ {para} ]]"
                    run = elem_p.add_run(note_text)
                    _set_run_font(run, size=Pt(10), italic=True, color=RGBColor(119, 119, 119))

        # Scene-end transition fallback:
        # If the scene does not already end with an explicit transition element, render scene.transition fallback.
        if not scene.is_sequence and not scene.has_scene_ending_transition(elements):
            scene_trans = (scene.transition or 'CUT TO').strip()
            if scene_trans:
                elem_p = doc.add_paragraph()
                elem_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elem_p.paragraph_format.space_before = Pt(12)
                elem_p.paragraph_format.space_after = Pt(12)
                run = elem_p.add_run(scene_trans.upper())
                _set_run_font(run, size=Pt(11), bold=True)

    # Apply custom watermark if configured on script
    watermark_text = (getattr(script, 'watermark', '') or '').strip()
    if watermark_text:
        _apply_docx_watermark(doc, watermark_text)

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    buffer.close()
    return docx_bytes

